from __future__ import annotations

import io

from app.models import DocumentoBruto
from app.readers.base import DocumentoIlegivel, ErroLeitura, LeitorArquivo, _normalizar_texto


class DocxReader(LeitorArquivo):
    formatos = ("docx",)

    def ler(self, dados: bytes, nome_arquivo: str) -> DocumentoBruto:
        try:
            import docx  # python-docx
        except ImportError as e:  # pragma: no cover
            raise ErroLeitura("python-docx não instalado (pip install python-docx).") from e

        try:
            doc = docx.Document(io.BytesIO(dados))
        except Exception as e:
            raise ErroLeitura(f"DOCX inválido: {e}") from e

        partes: list[str] = [p.text for p in doc.paragraphs]
        # tabelas costumam ter os campos de contrato/pedido
        for tabela in doc.tables:
            for linha in tabela.rows:
                celulas = [c.text.strip() for c in linha.cells]
                if any(celulas):
                    partes.append(" | ".join(celulas))

        texto = _normalizar_texto("\n".join(partes))
        if not texto:
            raise DocumentoIlegivel("DOCX sem texto.")
        return DocumentoBruto(texto=texto, paginas=1, metadados={"tabelas": len(doc.tables)})
